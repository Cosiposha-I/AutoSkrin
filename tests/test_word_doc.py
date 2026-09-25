"""Тесты записи скриншотов в документ Word (через python-docx, без Microsoft Word)."""

import os
from datetime import datetime

from docx import Document
from PIL import Image

import word_doc
from word_doc import ADDED_TO_FILE, PENDING, WordDocumentWriter, ensure_document


def image(color="red", size=(300, 200)):
    return Image.new("RGB", size, color)


def texts(path):
    return [p.text for p in Document(str(path)).paragraphs]


def test_creates_document_with_captions(tmp_path):
    path = tmp_path / "Конспект лекции.docx"
    writer = WordDocumentWriter(str(path))
    when = datetime(2024, 1, 15, 14, 30, 25)
    assert writer.add(image(), when) == ADDED_TO_FILE
    assert writer.add(image("blue"), when) == ADDED_TO_FILE
    writer.close()
    doc = Document(str(path))
    assert len(doc.inline_shapes) == 2
    lines = texts(path)
    assert lines[0] == "Конспект лекции"                       # заголовок нового документа
    assert "Скриншот 1 — 15.01.2024 14:30:25" in lines
    assert "Скриншот 2 — 15.01.2024 14:30:25" in lines
    assert not list((tmp_path).glob("*.tmp"))                  # временных файлов не осталось


def test_appends_to_existing_document_and_continues_numbering(tmp_path):
    path = tmp_path / "lecture.docx"
    doc = Document()
    doc.add_paragraph("Мои заметки до лекции")
    doc.save(str(path))
    first = WordDocumentWriter(str(path))
    first.add(image(), datetime.now())
    first.close()
    second = WordDocumentWriter(str(path), captions=True)
    second.add(image(), datetime.now())
    second.close()
    lines = texts(path)
    assert lines[0] == "Мои заметки до лекции"                 # содержимое пользователя на месте
    assert any(line.startswith("Скриншот 2 —") for line in lines)
    assert len(Document(str(path)).inline_shapes) == 2


def test_without_captions_and_width_fits_page(tmp_path):
    path = tmp_path / "wide.docx"
    writer = WordDocumentWriter(str(path), captions=False)
    writer.add(image(size=(4000, 500)), datetime.now())
    writer.close()
    doc = Document(str(path))
    section = doc.sections[-1]
    assert doc.inline_shapes[0].width <= section.page_width - section.left_margin - section.right_margin
    assert not any(t.startswith("Скриншот") for t in texts(path))


def test_uses_existing_image_file(tmp_path):
    shot = tmp_path / "shot.png"
    image().save(shot)
    writer = WordDocumentWriter(str(tmp_path / "doc.docx"))
    writer.add(image(), datetime.now(), image_path=shot)
    writer.close()
    assert shot.exists()                                       # файл из папки не удаляется


def test_busy_document_queues_shots(tmp_path, monkeypatch):
    path = tmp_path / "busy.docx"
    writer = WordDocumentWriter(str(path))
    writer.add(image(), datetime.now())

    real_replace = os.replace

    def locked(src, dst):
        if str(dst) == str(path):
            raise PermissionError(13, "Файл открыт в другой программе")
        return real_replace(src, dst)

    monkeypatch.setattr(word_doc.os, "replace", locked)
    assert writer.add(image("green"), datetime.now()) == PENDING
    assert writer.add(image("blue"), datetime.now()) == PENDING
    assert len(writer.pending) == 2
    assert len(Document(str(path)).inline_shapes) == 1        # файл не испорчен

    monkeypatch.setattr(word_doc.os, "replace", real_replace)  # документ закрыли
    assert writer.add(image("yellow"), datetime.now()) == ADDED_TO_FILE
    assert not writer.pending
    assert len(Document(str(path)).inline_shapes) == 4        # без дублей и потерь
    writer.close()


def test_close_rescues_temporary_images(tmp_path, monkeypatch):
    path = tmp_path / "Лекция.docx"
    writer = WordDocumentWriter(str(path))

    def always_locked(src, dst):
        raise PermissionError(13, "занят")

    monkeypatch.setattr(word_doc.os, "replace", always_locked)
    writer.add(image(), datetime.now())
    left, rescue_dir = writer.close()
    assert left == 1
    assert rescue_dir == tmp_path / "Лекция — не добавленные скриншоты"
    assert len(list(rescue_dir.glob("*.png"))) == 1


def test_ensure_document(tmp_path):
    path = ensure_document(tmp_path / "Новый" / "Конспект.docx")
    assert texts(path) == ["Конспект"]
    ensure_document(path)                                      # существующий не трогаем
    assert texts(path) == ["Конспект"]
