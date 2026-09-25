"""
word_doc.py — добавление скриншотов в документ Word (.docx), например
для конспекта лекции в реальном времени.

Каждый снимок дописывается в конец документа с подписью
«Скриншот N — 15.01.2024 14:30:25». Способы записи:

  1. Документ открыт в Microsoft Word (Windows) — снимок вставляется прямо
     в открытое окно через COM-автоматизацию. Конспект растёт на глазах,
     и рядом можно печатать свои заметки.
  2. Документ закрыт — файл дописывается библиотекой python-docx
     (создаётся, если его ещё нет).
  3. Файл занят другой программой — снимки ждут в очереди и добавляются,
     как только документ освободится. Ничего не теряется.

Объект WordDocumentWriter нужно создавать и использовать в одном потоке
(в потоке мониторинга): COM на Windows привязан к потоку.
"""

from __future__ import annotations

import logging
import os
import shutil
import sys
import tempfile
import time
import uuid
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from PIL import Image

log = logging.getLogger(__name__)

EMU_PER_PIXEL = 9525  # 1 пиксель при 96 dpi в единицах Word (EMU)

# Результаты добавления снимка
ADDED_TO_WORD = "word"   # вставлен в документ, открытый в Word
ADDED_TO_FILE = "file"   # дописан в файл .docx
PENDING = "pending"      # документ занят — снимок в очереди


class DocumentBusyError(Exception):
    """Документ сейчас нельзя изменить (занят другой программой)."""


@dataclass
class _Shot:
    image_path: Path
    when: datetime
    temporary: bool  # файл создан только для документа — удалить после вставки


def caption_text(number: int, when: datetime) -> str:
    return f"Скриншот {number} — {when:%d.%m.%Y %H:%M:%S}"


def ensure_document(path: str | Path) -> Path:
    """Создаёт пустой документ-конспект, если его ещё нет (для кнопки «Открыть»)."""
    path = Path(path)
    if not path.exists():
        from docx import Document
        path.parent.mkdir(parents=True, exist_ok=True)
        doc = Document()
        doc.add_heading(path.stem, level=1)
        doc.save(str(path))
    return path


class WordDocumentWriter:
    """Дописывает скриншоты в конец документа Word."""

    COM_RETRIES = 3  # Word может быть занят (пользователь печатает) — пробуем ещё раз

    def __init__(self, doc_path: str, captions: bool = True,
                 image_format: str = "png", jpeg_quality: int = 90):
        self.path = Path(doc_path)
        self.captions = captions
        self.image_format = image_format
        self.jpeg_quality = jpeg_quality
        self.pending: list[_Shot] = []
        self._cached_doc = None        # документ python-docx, пока файл не менялся извне
        self._cached_stat = None
        self._com = self._init_com()

    # ------------------------------------------------------------ публичные методы

    def add(self, image: Image.Image, when: datetime, image_path: Path | None = None) -> str:
        """
        Добавляет снимок в документ. image_path — уже сохранённый файл снимка
        (если сохранение в папку включено), иначе создаётся временный файл.
        Возвращает ADDED_TO_WORD, ADDED_TO_FILE или PENDING.
        """
        temporary = image_path is None
        if temporary:
            image_path = self._write_temp_image(image, when)
        self.pending.append(_Shot(Path(image_path), when, temporary))
        return self.flush()

    def flush(self) -> str:
        """Пытается добавить в документ все снимки из очереди."""
        if not self.pending:
            return ADDED_TO_FILE
        word_doc = self._find_open_in_word()
        try:
            if word_doc is not None:
                for shot in list(self.pending):
                    self._insert_via_word(word_doc, shot)
                    self._done(shot)
                return ADDED_TO_WORD
            self._append_to_file(list(self.pending))
            for shot in list(self.pending):
                self._done(shot)
            return ADDED_TO_FILE
        except DocumentBusyError as exc:
            log.info("Документ занят, снимков в очереди: %d (%s)", len(self.pending), exc)
            return PENDING

    def close(self) -> tuple[int, Path | None]:
        """
        Последняя попытка добавить очередь. Возвращает (сколько снимков так и
        не попало в документ, папку, куда перенесены их временные файлы).
        """
        self.flush()
        left = len(self.pending)
        rescue_dir = None
        temporary = [shot for shot in self.pending if shot.temporary and shot.image_path.exists()]
        if temporary:
            rescue_dir = self.path.parent / f"{self.path.stem} — не добавленные скриншоты"
            rescue_dir.mkdir(parents=True, exist_ok=True)
            for shot in temporary:
                shutil.move(str(shot.image_path), str(rescue_dir / shot.image_path.name))
        self.pending.clear()
        if self._com:
            try:
                import pythoncom
                pythoncom.CoUninitialize()
            except Exception:  # noqa: BLE001
                pass
            self._com = False
        return left, rescue_dir

    # ------------------------------------------------------------ вспомогательное

    def _write_temp_image(self, image: Image.Image, when: datetime) -> Path:
        folder = Path(tempfile.gettempdir()) / "AutoSkrin"
        folder.mkdir(parents=True, exist_ok=True)
        ext = "jpg" if self.image_format == "jpg" else "png"
        path = folder / f"screenshot_{when:%Y-%m-%d_%H-%M-%S}_{uuid.uuid4().hex[:6]}.{ext}"
        if ext == "jpg":
            image.convert("RGB").save(path, "JPEG", quality=int(self.jpeg_quality))
        else:
            image.save(path, "PNG", compress_level=3)
        return path

    def _done(self, shot: _Shot) -> None:
        self.pending.remove(shot)
        if shot.temporary:
            try:
                shot.image_path.unlink()
            except OSError:
                pass

    # ------------------------------------------------------------ python-docx

    def _append_to_file(self, shots: list[_Shot]) -> None:
        """Дописывает снимки в файл .docx (атомарно: временный файл + замена)."""
        from docx import Document
        from docx.shared import Emu, Pt, RGBColor

        self.path.parent.mkdir(parents=True, exist_ok=True)
        try:
            stat = self.path.stat() if self.path.exists() else None
            key = (stat.st_mtime_ns, stat.st_size) if stat else None
            if key is not None and key == self._cached_stat and self._cached_doc is not None:
                doc = self._cached_doc  # файл не менялся с нашей записи — не перечитываем
            elif stat is not None:
                doc = Document(str(self.path))
            else:
                doc = Document()
                doc.add_heading(self.path.stem, level=1)
        except PermissionError as exc:
            raise DocumentBusyError(str(exc)) from exc

        # Кэш сбрасываем заранее: если запись не удастся, в объекте останутся
        # недописанные снимки, и при повторе они бы задвоились
        self._cached_doc = self._cached_stat = None
        section = doc.sections[-1]
        max_width = section.page_width - section.left_margin - section.right_margin
        number = len(doc.inline_shapes)
        for shot in shots:
            number += 1
            if self.captions:
                run = doc.add_paragraph().add_run(caption_text(number, shot.when))
                run.italic = True
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
            with Image.open(shot.image_path) as img:
                width_px = img.width
            doc.add_picture(str(shot.image_path), width=Emu(min(max_width, width_px * EMU_PER_PIXEL)))

        fd, tmp = tempfile.mkstemp(prefix=".autoskrin-", suffix=".docx.tmp", dir=self.path.parent)
        os.close(fd)
        try:
            doc.save(tmp)
            os.replace(tmp, self.path)
        except PermissionError as exc:
            raise DocumentBusyError(str(exc)) from exc
        finally:
            if os.path.exists(tmp):
                os.remove(tmp)
        stat = self.path.stat()
        self._cached_doc, self._cached_stat = doc, (stat.st_mtime_ns, stat.st_size)

    # ------------------------------------------------------------ Microsoft Word (COM)

    @staticmethod
    def _init_com() -> bool:
        if sys.platform != "win32":
            return False
        try:
            import pythoncom
            pythoncom.CoInitialize()
            return True
        except Exception as exc:  # noqa: BLE001 — pywin32 нет: работаем только с файлом
            log.info("COM недоступен, запись только в файл: %s", exc)
            return False

    def _find_open_in_word(self):
        """Документ, если он открыт в запущенном Microsoft Word, иначе None."""
        if not self._com:
            return None
        try:
            import win32com.client
            word = win32com.client.GetActiveObject("Word.Application")
        except Exception:  # noqa: BLE001 — Word не запущен
            return None
        target = os.path.normcase(os.path.abspath(self.path))
        by_name = None
        try:
            for i in range(1, word.Documents.Count + 1):
                doc = word.Documents.Item(i)
                full_name = str(doc.FullName)
                if full_name.lower().startswith(("http://", "https://")):
                    # Документ из OneDrive: Word показывает веб-адрес — сравниваем по имени файла
                    if full_name.rsplit("/", 1)[-1].lower() == self.path.name.lower():
                        by_name = doc
                elif os.path.normcase(os.path.abspath(full_name)) == target:
                    return doc
        except Exception as exc:  # noqa: BLE001
            log.info("Не удалось получить список документов Word: %s", exc)
        return by_name

    def _insert_via_word(self, doc, shot: _Shot) -> None:
        for attempt in range(self.COM_RETRIES):
            try:
                self._insert_via_word_once(doc, shot)
                return
            except Exception as exc:  # noqa: BLE001
                if attempt == self.COM_RETRIES - 1:
                    raise DocumentBusyError(f"Word не принял снимок: {exc}") from exc
                time.sleep(0.5)

    @staticmethod
    def _new_last_paragraph(doc):
        """Пустой абзац в конце документа (создаётся, если последний абзац не пуст)."""
        last = doc.Paragraphs.Last.Range
        if str(last.Text).strip() or last.InlineShapes.Count:
            doc.Content.InsertParagraphAfter()
        return doc.Paragraphs.Last.Range

    def _insert_via_word_once(self, doc, shot: _Shot) -> None:
        number = doc.InlineShapes.Count + 1
        if self.captions:
            rng = self._new_last_paragraph(doc)
            text = caption_text(number, shot.when)
            rng.InsertBefore(text)
            caption = doc.Range(rng.Start, rng.Start + len(text))
            caption.Font.Italic = True
            caption.Font.Size = 9
        rng = self._new_last_paragraph(doc)
        # AddPicture(FileName, LinkToFile, SaveWithDocument, Range)
        shape = doc.InlineShapes.AddPicture(str(shot.image_path), False, True, rng)
        setup = doc.PageSetup
        max_width = setup.PageWidth - setup.LeftMargin - setup.RightMargin
        if 0 < max_width < shape.Width:
            height = shape.Height * max_width / shape.Width
            shape.LockAspectRatio = 0
            shape.Width = max_width
            shape.Height = height
        # Пустой абзац после снимка — чтобы заметки продолжались ниже картинки
        doc.Content.InsertParagraphAfter()
